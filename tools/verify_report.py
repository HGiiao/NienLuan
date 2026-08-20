# -*- coding: utf-8 -*-
import docx
DOC = r'D:\NienLuan\BaoCao_DoAn_HoanChinh.docx'
d = docx.Document(DOC)
paras = d.paragraphs
figs = [p.text for p in paras if p.text.strip().startswith('Hình ')]
tabs = [p.text for p in paras if p.text.strip().startswith('Bảng ')]
ph = [p.text for p in paras if p.text.strip().startswith('[CH')]
print('total_paragraphs:', len(paras))
print('figure_captions:', len(figs))
print('table_captions:', len(tabs))
print('remaining_placeholders:', len(ph))
print('inline_images:', len(d.inline_shapes))
# show structure placeholders area: find image paragraphs and following caption
for i, p in enumerate(paras):
    if len(p.runs) and 'graphicData' in p._p.xml:
        following = paras[i+1].text.strip() if i + 1 < len(paras) else '(none)'
        print('IMG# -> next para:', following[:70])
